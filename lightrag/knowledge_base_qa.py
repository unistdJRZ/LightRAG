from __future__ import annotations

import json
from io import BytesIO
from datetime import datetime, timezone
from typing import Any

from lightrag.utils import compute_mdhash_id, logger

_EXCEL_HEADER_ALIASES = {
    "kbqa_id": {
        "kbqa_id",
        "kbqa-id",
        "id",
        "qa_id",
        "qa-id",
        "编号",
        "问题编号",
    },
    "question": {
        "question",
        "preset_question",
        "预设问题",
        "问题",
    },
    "answer": {
        "answer",
        "preset_answer",
        "预设答案",
        "答案",
    },
    "metadata": {
        "metadata",
        "meta",
        "扩展信息",
        "元数据",
    },
}


def _normalize_excel_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _canonical_excel_header(value: Any) -> str | None:
    normalized = _normalize_excel_header(value)
    for canonical, aliases in _EXCEL_HEADER_ALIASES.items():
        if normalized in aliases:
            return canonical
    return None


def build_knowledge_base_qa_vector_data(
    qa_rows: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    """Build vector payloads keyed by knowledge-base QA id."""

    vector_data: dict[str, dict[str, Any]] = {}
    for row in qa_rows:
        if not isinstance(row, dict):
            continue
        kbqa_id = str(row.get("id") or row.get("kbqa_id") or "").strip()
        question = str(row.get("question") or "").strip()
        answer = str(row.get("answer") or "").strip()
        if not question or not answer:
            continue
        if not kbqa_id:
            kbqa_id = compute_mdhash_id(
                json.dumps(
                    {"question": question, "answer": answer},
                    ensure_ascii=False,
                    sort_keys=True,
                ),
                prefix="kbqa-",
            )
        vector_data[kbqa_id] = {
            "content": question,
            "kbqa_id": kbqa_id,
        }
    return vector_data


def parse_knowledge_base_qa_excel(file_bytes: bytes) -> list[dict[str, Any]]:
    """Parse a knowledge-base QA preset workbook.

    The first worksheet named `knowledge_base_qa` is preferred. If it does not
    exist, the active worksheet is used. Row 1 must contain headers. Required
    columns are `question` and `answer`; `kbqa_id` and `metadata` are optional.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency issue
        raise ValueError("openpyxl is required to import knowledge-base QA Excel") from exc

    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Invalid Excel workbook: {exc}") from exc

    worksheet = (
        workbook["knowledge_base_qa"]
        if "knowledge_base_qa" in workbook.sheetnames
        else workbook.active
    )
    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("Excel workbook is empty") from exc

    column_by_field: dict[str, int] = {}
    for index, header in enumerate(header_row):
        canonical = _canonical_excel_header(header)
        if canonical and canonical not in column_by_field:
            column_by_field[canonical] = index

    missing = [field for field in ("question", "answer") if field not in column_by_field]
    if missing:
        raise ValueError(
            "Excel is missing required column(s): " + ", ".join(missing)
        )

    parsed_rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for row_number, row in enumerate(rows_iter, start=2):
        values = list(row or [])
        if not any(value not in (None, "") for value in values):
            continue

        def cell(field: str) -> Any:
            index = column_by_field.get(field)
            if index is None or index >= len(values):
                return None
            return values[index]

        question = str(cell("question") or "").strip()
        answer = str(cell("answer") or "").strip()
        if not question:
            errors.append(f"row {row_number}: question is required")
        if not answer:
            errors.append(f"row {row_number}: answer is required")
        if not question or not answer:
            continue

        metadata: dict[str, Any] = {}
        metadata_value = cell("metadata")
        if metadata_value not in (None, ""):
            try:
                parsed_metadata = json.loads(str(metadata_value))
            except json.JSONDecodeError:
                errors.append(f"row {row_number}: metadata must be valid JSON")
                continue
            if not isinstance(parsed_metadata, dict):
                errors.append(f"row {row_number}: metadata must be a JSON object")
                continue
            metadata = parsed_metadata

        parsed_rows.append(
            {
                "id": str(cell("kbqa_id") or "").strip(),
                "question": question,
                "answer": answer,
                "metadata": metadata,
            }
        )

    if errors:
        raise ValueError("; ".join(errors))
    if not parsed_rows:
        raise ValueError("Excel workbook does not contain any valid QA rows")
    return parsed_rows


async def ensure_knowledge_base_qa_table(db: Any) -> None:
    await db.execute(
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_KNOWLEDGE_BASE_QA (
            workspace VARCHAR(255) NOT NULL,
            id VARCHAR(255) NOT NULL,
            question TEXT NOT NULL,
            answer TEXT NOT NULL,
            metadata JSONB NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMP(0) DEFAULT CURRENT_TIMESTAMP,
            update_time TIMESTAMP(0) DEFAULT CURRENT_TIMESTAMP,
            CONSTRAINT LIGHTRAG_KNOWLEDGE_BASE_QA_PK PRIMARY KEY (workspace, id)
        )
        """
    )
    await db.execute(
        "CREATE INDEX IF NOT EXISTS idx_lightrag_knowledge_base_qa_workspace ON LIGHTRAG_KNOWLEDGE_BASE_QA(workspace)"
    )


async def upsert_knowledge_base_qa_rows(
    *,
    db: Any,
    workspace: str,
    qa_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Insert or update knowledge-base QA rows and return normalized rows."""

    await ensure_knowledge_base_qa_table(db)
    normalized_rows: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc).replace(tzinfo=None)

    for row in qa_rows:
        if not isinstance(row, dict):
            continue
        question = str(row.get("question") or "").strip()
        answer = str(row.get("answer") or "").strip()
        if not question or not answer:
            continue
        kbqa_id = str(row.get("id") or row.get("kbqa_id") or "").strip()
        if not kbqa_id:
            kbqa_id = compute_mdhash_id(
                json.dumps(
                    {"workspace": workspace, "question": question, "answer": answer},
                    ensure_ascii=False,
                    sort_keys=True,
                ),
                prefix="kbqa-",
            )
        metadata = row.get("metadata") if isinstance(row.get("metadata"), dict) else {}
        await db.execute(
            """
            INSERT INTO LIGHTRAG_KNOWLEDGE_BASE_QA(
                workspace, id, question, answer, metadata, create_time, update_time
            )
            VALUES ($1, $2, $3, $4, $5, $6, $7)
            ON CONFLICT (workspace, id) DO UPDATE
            SET question=EXCLUDED.question,
                answer=EXCLUDED.answer,
                metadata=EXCLUDED.metadata,
                update_time=EXCLUDED.update_time
            """,
            {
                "workspace": workspace,
                "id": kbqa_id,
                "question": question,
                "answer": answer,
                "metadata": json.dumps(metadata, ensure_ascii=False),
                "create_time": now,
                "update_time": now,
            },
        )
        normalized_rows.append(
            {
                "id": kbqa_id,
                "question": question,
                "answer": answer,
                "metadata": metadata,
            }
        )

    logger.info("Stored %d knowledge-base QA pairs for workspace=%s", len(normalized_rows), workspace)
    return normalized_rows


async def get_knowledge_base_qa_by_ids(
    *,
    db: Any,
    workspace: str,
    kbqa_ids: list[str],
) -> list[dict[str, Any]]:
    if not kbqa_ids:
        return []

    await ensure_knowledge_base_qa_table(db)
    rows = await db.query(
        """
        SELECT id, workspace, question, answer, metadata,
               EXTRACT(EPOCH FROM create_time)::BIGINT AS created_at,
               EXTRACT(EPOCH FROM update_time)::BIGINT AS updated_at
        FROM LIGHTRAG_KNOWLEDGE_BASE_QA
        WHERE workspace=$1 AND id = ANY($2)
        """,
        [workspace, kbqa_ids],
        multirows=True,
    )
    rows_by_id = {str(row["id"]): row for row in rows or []}
    ordered_rows: list[dict[str, Any]] = []
    for kbqa_id in kbqa_ids:
        row = rows_by_id.get(str(kbqa_id))
        if not row:
            continue
        metadata = row.get("metadata")
        if isinstance(metadata, str):
            try:
                row["metadata"] = json.loads(metadata)
            except json.JSONDecodeError:
                row["metadata"] = {}
        ordered_rows.append(row)
    return ordered_rows


async def query_knowledge_base_qa_by_rule(
    *,
    db: Any,
    workspace: str,
    query: str,
    top_k: int,
) -> list[dict[str, Any]]:
    await ensure_knowledge_base_qa_table(db)
    query_text = str(query or "").strip()
    if not query_text:
        return []

    rows = await db.query(
        """
        SELECT id, workspace, question, answer, metadata,
               EXTRACT(EPOCH FROM create_time)::BIGINT AS created_at,
               EXTRACT(EPOCH FROM update_time)::BIGINT AS updated_at,
               CASE
                   WHEN lower(question) = lower($2) THEN 0
                   WHEN question ILIKE $3 THEN 1
                   WHEN $2 ILIKE '%' || question || '%' THEN 2
                   ELSE 3
               END AS match_rank
        FROM LIGHTRAG_KNOWLEDGE_BASE_QA
        WHERE workspace=$1
          AND (
              lower(question) = lower($2)
              OR question ILIKE $3
              OR $2 ILIKE '%' || question || '%'
          )
        ORDER BY match_rank ASC, update_time DESC
        LIMIT $4
        """,
        [workspace, query_text, f"%{query_text}%", top_k],
        multirows=True,
    )
    for row in rows or []:
        metadata = row.get("metadata")
        if isinstance(metadata, str):
            try:
                row["metadata"] = json.loads(metadata)
            except json.JSONDecodeError:
                row["metadata"] = {}
    return rows or []
